"""Blank paper organizers derived from the same definitions as browser support.

Only labels are consumed. Stored givens, statuses, and intermediate answers
must never be prefilled in the student's organizer.
"""
from docx.shared import Inches, Pt


def organizer_groups(problem):
    from utils.quantities import PROFILES, QUANTITIES
    for key in ("projectile_plan", "energy_plan"):
        if problem.get(key):
            plan = problem[key]
            return [("Given information", [f["label"] for f in plan["fields"]]),
                    ("Possible in-between calculations", [f["label"] for f in plan.get("steps", [])])]
    if problem.get("force_plan"):
        plan = problem["force_plan"]
        return [("Given information", list(plan["given"])),
                ("Possible in-between calculations", [s["label"] for s in plan["steps"]])]
    if problem.get("setup"):
        roles = PROFILES[problem["setup"]["profile"]]
        labels = [f"{role.name} ({QUANTITIES[role.quantity].unit})" for role in roles.values()]
        return [("Enter givens; ? for the target; x for unused quantities", labels)]
    return [("Plan your solution", ["Known information", "Find", "Relationship / plan"])]


def add_print_organizer(target, problem):
    """Compact labeled blanks; includes no solved values or automatic feedback."""
    heading = target.add_paragraph("Organizing help")
    heading.runs[0].bold = True
    heading.paragraph_format.space_before = Pt(2)
    heading.paragraph_format.space_after = Pt(2)
    for title, labels in organizer_groups(problem):
        if not labels:
            continue
        caption = target.add_paragraph(title)
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(2)
        caption.runs[0].font.size = Pt(9)
        table = target.add_table(rows=(len(labels)+2)//3, cols=3)
        table.style = "Table Grid"
        table.autofit = False
        for i, cell in enumerate(cell for row in table.rows for cell in row.cells):
            cell.width = Inches(2)
            if i < len(labels):
                cell.text = labels[i] + "\n\n"
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                for run in paragraph.runs:
                    run.font.size = Pt(9)
        # A Word cell needs a paragraph after each nested table. Its default
        # line height adds a blank line between organizer groups.
        spacer = target.paragraphs[-1]
        if not spacer.text:
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.line_spacing = Pt(1)

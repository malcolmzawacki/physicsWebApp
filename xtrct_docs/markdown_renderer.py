"""CommonMark blocks and inline formatting rendered as native Word content."""
import re
from urllib.parse import urlsplit
from markdown_it import MarkdownIt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt
from docx.table import _Cell
from xtrct_docs.word_math import office_math


def inline_math(state, silent):
    start = state.pos
    delimiter = next((d for d in ("$$", r"\(", r"\[", "$") if state.src.startswith(d, start)), None)
    if delimiter is None:
        return False
    close = {r"\(": r"\)", r"\[": r"\]"}.get(delimiter, delimiter)
    end = state.src.find(close, start + len(delimiter))
    while end >= 0 and state.src[end-1] == "\\":
        end = state.src.find(close, end + len(close))
    if end < 0:
        return False
    if not silent:
        token = state.push("math_inline", "", 0)
        token.content = state.src[start + len(delimiter):end]
    state.pos = end + len(close)
    return True


def block_math(state, start, end, silent):
    line = state.src[state.bMarks[start] + state.tShift[start]:state.eMarks[start]].strip()
    opening = next((x for x in ("$$", r"\[") if line.startswith(x)), None)
    if opening is None:
        return False
    closing = "$$" if opening == "$$" else r"\]"
    parts = [line[len(opening):]]
    last = start
    while closing not in parts[-1]:
        last += 1
        if last >= end:
            return False
        parts.append(state.src[state.bMarks[last] + state.tShift[last]:state.eMarks[last]])
    final, tail = parts[-1].split(closing, 1)
    if tail.strip():
        return False
    parts[-1] = final
    if not silent:
        token = state.push("math_block", "", 0)
        token.content = "\n".join(parts).strip()
        token.map = [start, last + 1]
    state.line = last + 1
    return True


def parser():
    md = MarkdownIt("commonmark", {"html": False}).enable(["table", "strikethrough"])
    md.inline.ruler.before("escape", "math_inline", inline_math)
    md.block.ruler.before("fence", "math_block", block_math, {"alt": ["paragraph", "reference", "blockquote", "list"]})
    return md


def render_inline(paragraph, tokens):
    bold = italic = strike = 0
    link = None
    for token in tokens:
        kind = token.type
        if kind == "strong_open": bold += 1
        elif kind == "strong_close": bold -= 1
        elif kind == "em_open": italic += 1
        elif kind == "em_close": italic -= 1
        elif kind == "s_open": strike += 1
        elif kind == "s_close": strike -= 1
        elif kind == "link_open":
            href = token.attrGet("href")
            if urlsplit(href).scheme not in ("http", "https", "mailto", ""):
                raise ValueError(f"Unsupported worksheet link: {href}")
            link = OxmlElement("w:hyperlink")
            if href.startswith("#"):
                link.set(qn("w:anchor"), href[1:])
            else:
                link.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
            paragraph._p.append(link)
        elif kind == "link_close": link = None
        elif kind == "math_inline":
            (paragraph._p if link is None else link).append(office_math(token.content))
        elif kind == "image":
            raise ValueError("Markdown images are not fetched; provide diagram_data and a renderer for worksheet images")
        elif kind in ("text", "code_inline", "softbreak", "hardbreak"):
            run = paragraph.add_run(" " if kind == "softbreak" else "" if kind == "hardbreak" else token.content)
            if kind == "hardbreak": run.add_break()
            run.bold, run.italic, run.font.strike = bool(bold), bool(italic), bool(strike)
            if kind == "code_inline": run.font.name = "Consolas"
            if link is not None:
                run.style = "Hyperlink" if "Hyperlink" in paragraph.part.document.styles else "Default Paragraph Font"
                run.underline = True
                link.append(run._r)
        else:
            raise ValueError(f"Unsupported inline worksheet markup: {kind}")


def add_inline(paragraph, text):
    render_inline(paragraph, parser().parseInline(text)[0].children or [])


def numbering(part, ordered, start):
    root = part.numbering_part.element
    abstracts = root.findall(qn("w:abstractNum"))
    abstract_id = max((int(n.get(qn("w:abstractNumId"))) for n in abstracts), default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl"); level.set(qn("w:ilvl"), "0")
    for name, value in (("start", start), ("numFmt", "decimal" if ordered else "bullet"), ("lvlText", "%1." if ordered else "•")):
        node = OxmlElement("w:" + name); node.set(qn("w:val"), str(value)); level.append(node)
    abstract.append(level)
    # Abstract definitions precede concrete numbering instances in Word's schema.
    root.insert(len(abstracts), abstract)
    instance = root.add_num(abstract_id)
    return instance.numId


def add_markdown(cell, text, first=None):
    previous = {p._p for p in cell.paragraphs}
    tokens = parser().parse(text)
    paragraph = first
    available_first = first
    lists = []
    item_pending = False
    quote_depth = 0
    table = row = None
    header = False
    cell_index = 0

    def new_paragraph(style=None):
        nonlocal available_first, item_pending
        current = available_first if available_first is not None else cell.add_paragraph()
        available_first = None
        if style: current.style = style
        if lists:
            current.paragraph_format.left_indent = Inches(.25 * len(lists))
            if item_pending:
                num = current._p.get_or_add_pPr().get_or_add_numPr()
                num.get_or_add_ilvl().val = 0
                num.get_or_add_numId().val = lists[-1]
                current.paragraph_format.first_line_indent = Inches(-.18)
                item_pending = False
        if quote_depth:
            current.paragraph_format.left_indent = Inches(.25 * (len(lists) + quote_depth))
        return current

    for token in tokens:
        kind = token.type
        if kind in ("heading_open", "bullet_list_open", "ordered_list_open", "blockquote_open", "fence", "code_block", "math_block", "table_open"):
            available_first = None
        if kind in ("bullet_list_open", "ordered_list_open"):
            lists.append(numbering(cell.part, kind == "ordered_list_open", token.attrGet("start") or 1))
        elif kind in ("bullet_list_close", "ordered_list_close"): lists.pop()
        elif kind == "list_item_open": item_pending = True
        elif kind == "list_item_close": item_pending = False
        elif kind == "blockquote_open": quote_depth += 1
        elif kind == "blockquote_close": quote_depth -= 1
        elif kind in ("paragraph_open", "heading_open"):
            paragraph = new_paragraph("Heading " + str(min(9, int(token.tag[1:]) + 1)) if kind == "heading_open" else None)
        elif kind == "inline":
            if paragraph is None: paragraph = new_paragraph()
            render_inline(paragraph, token.children or [])
        elif kind in ("fence", "code_block"):
            paragraph = new_paragraph()
            run = paragraph.add_run(token.content.rstrip("\n")); run.font.name = "Consolas"; run.font.size = Pt(10)
        elif kind == "math_block":
            paragraph = new_paragraph()
            paragraph.alignment = 1
            paragraph._p.append(office_math(token.content))
        elif kind == "table_open":
            available_first = None
            table_width = (cell.width or Inches(5.5)) if isinstance(cell, _Cell) else (
                cell.sections[-1].page_width - cell.sections[-1].left_margin - cell.sections[-1].right_margin)
            table = cell.add_table(rows=0, cols=0)
            table.style = "Table Grid"
        elif kind == "thead_open": header = True
        elif kind == "thead_close": header = False
        elif kind == "tr_open":
            row = table.add_row(); cell_index = 0
            if header:
                repeat = OxmlElement("w:tblHeader"); row._tr.get_or_add_trPr().append(repeat)
        elif kind in ("td_open", "th_open"):
            if cell_index >= len(table.columns): table.add_column(Inches(1))
            paragraph = row.cells[cell_index].paragraphs[0]
            style = token.attrGet("style") or ""
            paragraph.alignment = 2 if "right" in style else 1 if "center" in style else 0
            if header:
                shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "E6E6E6"); row.cells[cell_index]._tc.get_or_add_tcPr().append(shading)
            cell_index += 1
        elif kind == "table_close":
            width = int(table_width / max(1, len(table.columns)))
            for column in table.columns: column.width = width
            for row in table.rows:
                for table_cell in row.cells: table_cell.width = width
            paragraph = cell.add_paragraph(); table = None
        elif kind == "hr":
            paragraph = new_paragraph()
            border = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
            for key, value in (("val", "single"), ("sz", "4"), ("color", "CCCCCC")): bottom.set(qn("w:"+key), value)
            border.append(bottom); paragraph._p.get_or_add_pPr().append(border)
        elif kind.endswith("_close") or kind == "tbody_open":
            pass
        else:
            raise ValueError(f"Unsupported worksheet block: {kind}")
    for p in cell.paragraphs:
        if p._p not in previous or (first is not None and p._p is first._p):
            p.paragraph_format.keep_together = True
            p.paragraph_format.keep_with_next = True
    return cell.paragraphs[-1]


def add_question(cell, text, number):
    first = cell.paragraphs[0]
    first.add_run(f"{number}. ")
    # First-line indentation is often an f-string artifact, not a code block.
    lines = text.strip().splitlines()
    if len(lines) > 1:
        # Remove a consistent continuation indent without destroying real lists.
        nonblank = [line for line in lines[1:] if line.strip()]
        indent = min((len(line) - len(line.lstrip()) for line in nonblank), default=0)
        if indent >= 8:
            lines = [lines[0], *[line[indent:] if line.strip() else "" for line in lines[1:]]]
    return add_markdown(cell, "\n".join(lines), first)

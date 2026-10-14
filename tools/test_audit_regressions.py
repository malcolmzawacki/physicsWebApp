"""Physics, export and interaction regressions from the September audit."""
import io
import logging
import math
import random
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
logging.disable(logging.CRITICAL)
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from streamlit.testing.v1 import AppTest
from utils.problem_payload import payload_from_dict, ProblemPayloadError


def page(section, activity, group=None):
    app = AppTest.from_file(str(ROOT / "Home.py"), default_timeout=25)
    app.session_state["nav_level"] = "advanced"
    app.session_state["router_section"] = section
    app.session_state["router_group"] = group
    app.session_state["router_activity"] = activity
    return app.run()


class AuditRegressions(unittest.TestCase):
    def tearDown(self):
        plt.close("all")

    def test_tension_equilibrium(self):
        from utils.generators.forces.tension_generator import TensionGenerator
        generator = TensionGenerator()
        for difficulty in ("Easy", "Medium", "Hard"):
            for seed in range(500):
                random.seed(seed)
                case = generator.tension_scenario(difficulty)
                force = [0., 0.]
                for i, quadrant in enumerate(case["scenario"], 1):
                    tension = case[f"tension {i}"]
                    self.assertGreater(tension, 0)
                    angle = math.radians(case[f"theta {i}"])
                    sx, sy = generator.quadrant_phrases(quadrant)["signs"]
                    force[0] += sx * tension * math.cos(angle)
                    force[1] += sy * tension * math.sin(angle)
                self.assertAlmostEqual(force[0], 0, places=7)
                self.assertAlmostEqual(force[1], case["mass"] * 10, places=7)

    def test_thermal_domain(self):
        from utils.generators.energy.thermal_loss import ThermalLossGenerator, sample_thermal_loss
        self.assertGreater(sample_thermal_loss(1), 0)
        self.assertLess(sample_thermal_loss(1), 1)
        generator = ThermalLossGenerator()
        for seed in range(1000):
            random.seed(seed)
            for kind in generator.stored_metadata():
                result = generator.choose_problem_dict(kind, "Hard")
                payload_from_dict(result)
                self.assertTrue(all(a > 0 for a in result["answers"]), (seed, result))

    def test_parallel_connectivity(self):
        import schemdraw
        import schemdraw.elements as elm
        from utils.circuit_diagrams import draw_parallel
        drawing = schemdraw.Drawing(show=False)
        draw_parallel(drawing, elm, {"source_label": "12 V", "resistor_labels": ["R1", "R2"]})
        node = lambda point: tuple(round(v, 7) for v in point)
        battery = drawing.elements[0]
        wires = {}
        resistors = []
        for element in drawing.elements[1:]:
            a, b = node(element.start), node(element.end)
            if isinstance(element, elm.Line):
                wires.setdefault(a, set()).add(b)
                wires.setdefault(b, set()).add(a)
            else:
                resistors.append((a, b))
        def connected(start):
            seen, pending = set(), [start]
            while pending:
                n = pending.pop()
                if n not in seen:
                    seen.add(n)
                    pending.extend(wires.get(n, ()))
            return seen
        top, bottom = connected(node(battery.end)), connected(node(battery.start))
        self.assertFalse(top & bottom, "A wire short-circuits the source")
        self.assertEqual(len(resistors), 2)
        self.assertTrue(all(a in top and b in bottom for a, b in resistors))

    def test_circuit_export_survives_another_pages_dark_theme(self):
        from utils.circuit_diagrams import circuit_figure
        from PIL import Image, ImageStat
        with plt.style.context("dark_background"):
            figure = circuit_figure({"kind": "parallel", "source_label": "12 V",
                "resistor_labels": ["R1 = 6 ohms", "R2 = 3 ohms"]})
            image = io.BytesIO()
            figure.savefig(image, format="png", bbox_inches="tight")
            image.seek(0)
            gray = Image.open(image).convert("L")
            self.assertLess(gray.getextrema()[0], 20)
            self.assertGreater(ImageStat.Stat(gray).mean[0], 180)
            self.assertEqual(figure.axes[0].get_xlim()[0] < 0, True)

    def test_payload_ownership_and_numeric_domain(self):
        source = {"question": "Find x", "answers": [1], "units": ["m"], "extras": {"case": 7}, "explanation": "Worked solution"}
        normalized = payload_from_dict(source)
        self.assertEqual(source["extras"], {"case": 7})
        self.assertEqual(normalized.extras["explanation"], "Worked solution")
        for answer in (complex(0, 1), float("nan"), float("inf")):
            with self.assertRaises(ProblemPayloadError):
                payload_from_dict({**source, "answers": [answer]})

    def test_stable_progress_identity_migrates_legacy_labels(self):
        app = AppTest.from_string('''
from utils.ui import Interface
class Generator:
    def get_problem_metadata(self, kind):
        return {"id": "constant-speed", "aliases": ["Old label"]}
ui = Interface("identity", "Practice", Generator(), ["New label"], ["Medium"])
ui.initialize_session_state()
ui.performance_dropdown()
''')
        app.session_state["identity_performance"] = {"Old label": {"Medium": {"attempts": 4, "correct": 3}}}
        app.run()
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state["identity_performance"], {"constant-speed": {"Medium": {"attempts": 4, "correct": 3}}})

    def test_graph_difficulty_and_ranking_constraint(self):
        from utils.generators.kinematics.motion_graph_generator import MotionGraphGenerator
        from utils.generators.torque_generator import TorqueGenerator
        generator = MotionGraphGenerator()
        for seed in range(10):
            random.seed(seed)
            self.assertTrue(generator.choose_problem_dict("Position-Time Graph", "Easy")["graph_type"].startswith("linear"))
        # Every ranking payload must explicitly carry the uniqueness constraint.
        torque = TorqueGenerator()
        for kind in torque.stored_metadata():
            result = torque.choose_problem_dict(kind, "Hard")
            if result.get("answer_input_mode") == "dropdown":
                self.assertTrue(result["unique_choices"])

    def test_every_worksheet_factory_serializes(self):
        from xtrct_docs.question_organizer import Org
        from xtrct_docs.document_creator import create_doc
        org = Org()
        names = ("funny_forces_practice", "projectile_quiz_CP", "projectile_quiz_H", "create_first_doc",
                 "unit1_practice_doc", "constant_motion_quiz", "projectile_practice", "motion_graph_test",
                 "mixed_motion_with_graphs", "mixed_motion_with_graphs_H")
        for name in names:
            with self.subTest(factory=name):
                random.seed(2)
                factory = getattr(org, name)()
                self.assertTrue(callable(factory))
                output = io.BytesIO()
                create_doc("Regression worksheet", factory, 1, include_graphs=False, output_path=output, open_document=False)
                output.seek(0)
                document = Document(output)
                self.assertGreater(len(document.tables), 0)
                self.assertEqual(len(document.inline_shapes), 0)
                plt.close("all")

    def test_diagram_export_and_prompt_structure(self):
        from xtrct_docs.question_organizer import Org
        from xtrct_docs.document_creator import create_doc
        from utils.problem_adapters import circuit_payload, matching_payload
        from utils.generators.current_electricity import build_series_parallel_case, build_circuit_case
        from utils.generators.kinematics.motion_graph_matching_generator import generate_motion_graph_match_payload
        org = Org()
        problems = [org.TG.suspension_question("Hard"), org.AWG.static_half_atwood(),
                    org.DDG.distance_and_displacement_1D("Easy"),
                    circuit_payload(build_series_parallel_case("Parallel Branch Current", "Hard")),
                    circuit_payload(build_circuit_case("Series Current and Voltage", "Hard")),
                    circuit_payload(build_circuit_case("Parallel Current and Voltage", "Hard")),
                    matching_payload(generate_motion_graph_match_payload("Position-Time First")),
                    {"question": "Find **speed**.\n\n- Distance: 10 m\n- Time: 2 s\n\nUse $v=\\frac{d}{t}$.",
                     "answers": [5, "East"], "units": ["m/s", "Direction"], "button_options": {1: ["East", "West"]}}]
        output = io.BytesIO()
        factory = lambda: [{"heading": "Mixed practice", "gap": 0, "problems": problems}]
        create_doc("Diagram regression", factory, 1, output_path=output, open_document=False)
        output.seek(0)
        doc = Document(output)
        self.assertGreaterEqual(len(doc.inline_shapes), 7)
        self.assertTrue(doc.element.xpath("//m:oMath/m:f"))
        last = doc.tables[-1].cell(0, 0)
        texts = [p.text for p in last.paragraphs]
        self.assertTrue(any("Distance: 10 m" in text for text in texts))
        self.assertTrue(any(r.bold and "speed" in r.text for p in last.paragraphs for r in p.runs))
        self.assertNotIn("**", " ".join(texts))

    def test_standard_pause_resume_and_invalid_numbers(self):
        app = page("Kinematics", "Constant Motion")
        prefix = "const_motion_"
        qid = app.session_state[prefix + "_question_id"]
        next(i for i in app.text_input if "setup_input" not in i.key).set_value("nan")
        next(b for b in app.button if b.label == "Submit").click().run()
        self.assertFalse(app.session_state[prefix + "_submitted"])
        app.session_state[f"{prefix}_cancel_next_{qid}"] = True
        final_inputs = [i for i in app.text_input if "setup_input" not in i.key]
        for field, answer in zip(final_inputs, app.session_state[prefix + "_correct_answers"]):
            field.set_value(str(answer))
        next(b for b in app.button if b.label == "Submit").click().run()
        self.assertTrue(app.success)
        app.run()
        self.assertTrue(app.success)
        self.assertEqual(app.session_state[prefix + "_stars"], 1)
        app.checkbox(key=f"{prefix}_cancel_next_{qid}").uncheck().run()
        self.assertEqual(app.session_state[prefix + "_question_id"], qid + 1)
        self.assertFalse(app.exception)

    def test_choice_reset_and_matching_single_score(self):
        app = page("Kinematics", "Types of Motion Graphs")
        next(b for b in app.button if "_option_0_" in b.key).click().run()
        app.button(key="motion_graph_new_question").click().run()
        self.assertEqual(app.session_state["motion_graph_user_answers_selected"], [None, None])
        app = page("Kinematics", "Matching Motion Graphs")
        qid = app.session_state["motion_graph_matching_question_id"]
        payload = app.session_state["motion_graph_matching_payload"]
        labels = payload.get("choice_labels") or [chr(65+i) for i in range(len(payload["option_diagrams"]))]
        app.session_state["motion_graph_matching_cards"] = {"question_id": qid, "index": payload["correct_index"]}
        app.run()
        app.session_state[f"motion_graph_matching_cancel_next_{qid}"] = True
        app.button(key="motion_graph_matching_check").click().run()
        app.run()
        self.assertEqual(app.session_state["motion_graph_matching_stars"], 1)
        self.assertTrue(app.button(key="motion_graph_matching_check").disabled)
        app.radio(key="motion_graph_matching_primary_select").set_value("Velocity-Time First").run()
        self.assertEqual(app.session_state["motion_graph_matching_question_id"], qid + 1)
        self.assertFalse(app.session_state["motion_graph_matching_submitted"])

    def test_matching_selection_layout_and_stale_events(self):
        from utils.graph_cards import selected_index
        for event in (None, {}, {"question_id": 1, "index": 0},
                      {"question_id": 2, "index": True},
                      {"question_id": 2, "index": -1},
                      {"question_id": 2, "index": 3}):
            self.assertIsNone(selected_index(event, 2, 3))
        self.assertEqual(selected_index({"question_id": 2, "index": 0}, 2, 3), 0)
        app = page("Kinematics", "Matching Motion Graphs")
        prefix = "motion_graph_matching_"
        qid = app.session_state[prefix + "question_id"]
        wrong = (app.session_state[prefix + "payload"]["correct_index"] + 1) % 3
        app.session_state[prefix + "cards"] = {"question_id": qid, "index": wrong}
        app.run()
        self.assertFalse(app.session_state[prefix + "submitted"])
        app.selectbox(key=prefix + "layout").select("Options in a column").run()
        self.assertEqual(app.session_state[prefix + "question_id"], qid)
        self.assertEqual(app.session_state[prefix + "selection"], wrong)
        app.button(key=prefix + "check").click().run()
        self.assertTrue(app.session_state[prefix + "submitted"])
        self.assertFalse(app.session_state[prefix + "last_result"])
        app.session_state[prefix + "cards"] = {"question_id": qid, "index": (wrong + 1) % 3}
        app.run()
        self.assertEqual(app.session_state[prefix + "selection"], wrong)
        app.button(key=prefix + "generate").click().run()
        self.assertEqual(app.session_state[prefix + "question_id"], qid + 1)
        self.assertIsNone(app.session_state[prefix + "selection"])
        self.assertFalse(app.exception)

    def test_mixed_inputs_and_duplicate_dropdowns(self):
        script = '''
from utils.ui import Interface
class Generator:
    def choose_problem_dict(self, kind, difficulty):
        return dict(question="Find speed and directions", answers=[5, "East", "East"],
            units=["m/s", "Direction 1", "Direction 2"], button_options={1: ["East", "West"], 2: ["East", "West"]},
            extras=dict(answer_input_mode="dropdown", explanation="Distance divided by time", time_limit=0))
    def get_problem_metadata(self, kind):
        return {"honors": ["v=d/t", "d=vt"]}
ui = Interface("mixed", "Mixed", Generator(), ["Practice"], ["Medium"])
ui.unified_smart_layout()
'''
        app = AppTest.from_string(script).run()
        self.assertFalse(app.exception)
        self.assertIn("gathered", app.latex[0].value)
        app.session_state["mixed_cancel_next_1"] = True
        app.text_input[0].set_value("5")
        app.selectbox(key="mixed_dropdown_1_1").select("East")
        app.selectbox(key="mixed_dropdown_2_1").select("East")
        next(b for b in app.button if b.label == "Submit Answers").click().run()
        self.assertTrue(app.success)
        self.assertEqual(app.session_state["mixed_stars"], 2)
        self.assertTrue(any("Distance divided by time" in m.value for m in app.markdown))
        app.checkbox(key="mixed_cancel_next_1").uncheck().run()
        self.assertEqual(app.session_state["mixed_question_id"], 2)

    def test_circuit_learning_focus_cases(self):
        from utils.generators.current_electricity import (
            PROBLEM_TYPES, SERIES_PARALLEL_PROBLEM_TYPES,
            build_circuit_case, build_series_parallel_case,
        )
        from utils.problem_adapters import circuit_payload
        self.assertNotIn("Series Current", PROBLEM_TYPES)
        self.assertNotIn("Voltage Drop", PROBLEM_TYPES)
        self.assertIn("Series Voltage Drop", SERIES_PARALLEL_PROBLEM_TYPES)
        for difficulty in ("Easy", "Medium", "Hard"):
            for seed in range(120):
                for kind in PROBLEM_TYPES[1:]:
                    random.seed(seed)
                    case = build_circuit_case(kind, difficulty)
                    givens = dict(case["givens"])
                    r1, r2 = (float(givens[k].split()[0]) for k in ("R1", "R2"))
                    if kind == "Series Current and Voltage":
                        current = float(givens["Shared current"].split()[0])
                        expected = [current * r1, current * r2]
                        self.assertNotIn("Battery voltage", givens)
                        self.assertEqual(case["unit"], "V")
                    else:
                        voltage = float(givens["Source voltage"].split()[0])
                        expected = [voltage / r1, voltage / r2]
                        self.assertEqual(case["diagram"]["kind"], "parallel")
                        self.assertEqual(case["unit"], "A")
                    self.assertEqual(len(case["answer_labels"]), 2)
                    for answer, value in zip(case["answers"], expected):
                        self.assertAlmostEqual(answer, value)
                    self.assertNotIn("setup", case)
                    payload = payload_from_dict(circuit_payload(case))
                    self.assertEqual(list(payload.answers), expected)
                    for resistance, label in zip((r1, r2), case["diagram"]["resistor_labels"]):
                        self.assertEqual(float(label.split("=")[1].split()[0]), resistance)
                random.seed(seed)
                case = build_series_parallel_case("Series Voltage Drop", difficulty)
                givens = dict(case["givens"])
                r1, r2 = (float(givens[k].split()[0]) for k in ("R1", "R2"))
                source = float(givens["Battery voltage"].split()[0])
                target = case["answer_label"].split()[-1]
                expected = source / (r1 + r2) * float(givens[target].split()[0])
                self.assertAlmostEqual(case["answer"], expected)
                self.assertNotIn("Shared current", givens)
                self.assertEqual(case["diagram"]["wire_label"], "")
                payload_from_dict(circuit_payload(case))
                for kind in ("Series Equivalent Resistance", "Parallel Equivalent Resistance"):
                    self.assertNotIn("?", build_series_parallel_case(kind, difficulty)["diagram"]["source_label"])

    def test_circuit_focus_routes_and_retired_selection(self):
        app = page("⚡Electricity⚡", "Circuit Ohm's Law", "Current Electricity")
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state["router_activity"], "Circuit Fundamentals")
        prefix = "current_electricity"
        app = AppTest.from_file(str(ROOT / "Home.py"), default_timeout=25)
        app.session_state["nav_level"] = "advanced"
        app.session_state["router_section"] = "⚡Electricity⚡"
        app.session_state["router_group"] = "Current Electricity"
        app.session_state["router_activity"] = "Circuit Fundamentals"
        app.session_state[prefix + "_problem_type"] = "Series Current"
        app.session_state[prefix + "_problem_type_select"] = "Series Current"
        old_qid = 4
        app.session_state[prefix + "_question_number"] = old_qid
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state[prefix + "_case"]["problem_type"], "Single Resistor")
        self.assertEqual(app.session_state[prefix + "_question_number"], old_qid + 1)
        for kind in ("Series Current and Voltage", "Parallel Current and Voltage"):
            app.selectbox(key=prefix + "_fundamentals_practice_select").select(kind).run()
            for difficulty in ("Easy", "Medium", "Hard"):
                app.selectbox(key=prefix + "_difficulty_select").select(difficulty).run()
                self.assertFalse(app.exception)
                case = app.session_state[prefix + "_case"]
                self.assertEqual(case["problem_type"], kind)
                self.assertTrue(any(e.label == "Organizing help" for e in app.expander))
                fields = [i for i in app.text_input if "setup_input" not in i.key]
                self.assertEqual([i.label for i in fields], [f"{label} ({case['unit']})" for label in case["answer_labels"]])
                before = app.session_state[prefix + "_attempt_count"]
                next(b for b in app.button if b.label == "Check Answer").click().run()
                self.assertEqual(app.session_state[prefix + "_attempt_count"], before)
                self.assertFalse(app.session_state[prefix + "_submitted"])
                fields = [i for i in app.text_input if "setup_input" not in i.key]
                fields[0].set_value(str(case["answers"][0]))
                fields[1].set_value("inf")
                next(b for b in app.button if b.label == "Check Answer").click().run()
                self.assertEqual(app.session_state[prefix + "_attempt_count"], before)
                self.assertFalse(app.session_state[prefix + "_submitted"])
                app.text_input(key=fields[1].key).set_value("-1")
                expected_results = [True, False]
                if difficulty == "Medium":
                    app.text_input(key=fields[0].key).set_value("-1")
                    app.text_input(key=fields[1].key).set_value(str(case["answers"][1]))
                    expected_results = [False, True]
                next(b for b in app.button if b.label == "Check Answer").click().run()
                self.assertEqual(app.session_state[prefix + "_last_result"]["field_results"], expected_results)
                self.assertTrue(app.error)
                self.assertTrue(all(i.disabled for i in app.text_input if "setup_input" not in i.key))
                self.assertFalse(any(c.label == "Cancel Next Question" for c in app.checkbox))
                self.assertTrue(any(i.value == case["explanation"] for i in app.info))
                app.run()
                self.assertEqual(app.session_state[prefix + "_attempt_count"], before + 1)
                next(b for b in app.button if b.label == "New Question (Circuit)").click().run()
                qid = app.session_state[prefix + "_question_number"]
                app.session_state[f"{prefix}_cancel_next_{qid}"] = True
                case = app.session_state[prefix + "_case"]
                fields = [i for i in app.text_input if "setup_input" not in i.key]
                self.assertTrue(all(not i.value for i in fields))
                for field, answer in zip(fields, case["answers"]):
                    field.set_value(str(round(answer, 2)))
                correct_before = app.session_state[prefix + "_correct_count"]
                next(b for b in app.button if b.label == "Check Answer").click().run()
                self.assertFalse(app.exception)
                self.assertTrue(app.success)
                self.assertEqual(app.session_state[prefix + "_last_result"]["field_results"], [True, True])
                self.assertEqual(app.session_state[prefix + "_attempt_count"], before + 2)
                self.assertEqual(app.session_state[prefix + "_correct_count"], correct_before + 1)
                app.run()
                self.assertEqual(app.session_state[prefix + "_attempt_count"], before + 2)
                self.assertEqual(app.session_state[prefix + "_question_number"], qid)
                app.session_state[prefix + "_advance_deadline"] = 0.0
                app.checkbox(key=f"{prefix}_cancel_next_{qid}").uncheck().run()
                self.assertEqual(app.session_state[prefix + "_question_number"], qid + 1)
                self.assertFalse(app.session_state[prefix + "_submitted"])
                app.run()
        app = page("⚡Electricity⚡", "Series and Parallel Circuits", "Current Electricity")
        prefix = "current_electricity_series_parallel"
        app.selectbox(key=prefix + "_problem_type_select").select("Series Voltage Drop").run()
        for difficulty in ("Easy", "Medium", "Hard"):
            app.selectbox(key=prefix + "_difficulty_select").select(difficulty).run()
            qid = app.session_state[prefix + "_question_number"]
            app.session_state[f"{prefix}_cancel_next_{qid}"] = True
            next(i for i in app.text_input if "setup_input" not in i.key).set_value(str(app.session_state[prefix + "_case"]["answer"]))
            next(b for b in app.button if b.label == "Check Answer").click().run()
            self.assertFalse(app.exception)
            self.assertTrue(app.success)
            self.assertEqual(app.session_state[prefix + "_question_number"], qid)

    def test_circuit_attempts_and_persistent_feedback(self):
        for activity, prefix in (("Circuit Fundamentals", "current_electricity"),
                ("Series and Parallel Circuits", "current_electricity_series_parallel")):
            app = page("⚡Electricity⚡", activity, "Current Electricity")
            self.assertFalse(any("#### Givens" in m.value for m in app.markdown))
            self.assertTrue(app.get("imgs"))
            self.assertEqual(sum(b.label == "New Question (Circuit)" for b in app.button), 1)
            self.assertFalse(any(b.label in ("New Circuit", "Next Circuit") for b in app.button))
            self.assertTrue(any(e.label == "Your Performance" for e in app.expander))
            qid = app.session_state[prefix + "_question_number"]
            next(i for i in app.text_input if "setup_input" not in i.key).set_value("inf")
            next(b for b in app.button if b.label == "Check Answer").click().run()
            self.assertEqual(app.session_state[prefix + "_attempt_count"], 0)
            app.session_state[f"{prefix}_cancel_next_{qid}"] = True
            next(i for i in app.text_input if "setup_input" not in i.key).set_value(str(app.session_state[prefix + "_case"]["answer"]))
            next(b for b in app.button if b.label == "Check Answer").click().run()
            app.run()
            self.assertEqual(app.session_state[prefix + "_attempt_count"], 1)
            self.assertTrue(app.success)
            self.assertEqual(app.session_state[prefix + "_question_number"], qid)
            calculation = next(e for e in app.expander if e.label == "View calculation")
            self.assertFalse(calculation.proto.expanded)
            self.assertTrue(any(app.session_state[prefix + "_case"]["explanation"] == i.value for i in app.info))
            self.assertTrue(app.checkbox(key=f"{prefix}_cancel_next_{qid}").value)
            message = app.session_state[prefix + "_last_result"]["message"]
            self.assertTrue(any(i.value == message for i in app.success))
            app.run()
            self.assertEqual(app.session_state[prefix + "_last_result"]["message"], message)
            app.session_state[prefix + "_advance_deadline"] = 0.0
            app.checkbox(key=f"{prefix}_cancel_next_{qid}").uncheck().run()
            self.assertEqual(app.session_state[prefix + "_question_number"], qid + 1)
            app.run()
            self.assertFalse(any(e.label == "View calculation" for e in app.expander))
            qid += 1
            next(i for i in app.text_input if "setup_input" not in i.key).set_value("-1")
            next(b for b in app.button if b.label == "Check Answer").click().run()
            from utils.word_lists import error_messages
            self.assertIn(app.session_state[prefix + "_last_result"]["message"], error_messages)
            self.assertFalse(any(c.label == "Cancel Next Question" for c in app.checkbox))
            explanation = app.session_state[prefix + "_case"]["explanation"]
            self.assertTrue(any(i.value == explanation for i in app.info))
            self.assertFalse(any(e.label == "View calculation" for e in app.expander))
            app.run()
            self.assertEqual(app.session_state[prefix + "_question_number"], qid)
            self.assertEqual(app.session_state[prefix + "_attempt_count"], 2)
            self.assertTrue(any(i.value == explanation for i in app.info))
            next(b for b in app.button if b.label == "New Question (Circuit)").click().run()
            self.assertEqual(app.session_state[prefix + "_question_number"], qid + 1)
            self.assertFalse(any(i.value == explanation for i in app.info))


if __name__ == "__main__":
    unittest.main()

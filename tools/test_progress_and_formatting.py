"""Stable history and native Word markup integration regressions."""
import io
import logging
import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
logging.disable(logging.CRITICAL)
from docx import Document
from docx.oxml.ns import qn
from streamlit.testing.v1 import AppTest
from tools.validate_payloads import discover_generators
from utils.progress_ids import CUSTOM_PROGRESS, migrate_custom_progress
from xtrct_docs.document_creator import create_doc
from xtrct_docs.word_math import office_math

FORMATTING_SAMPLE = r'''## Relationships and units

Use **bold with *nested emphasis***, ~~a correction~~, `literal $x$`, and a [reference](https://example.com/physics).
This line ends with a hard break.  
This is the next line. Escaped currency: \$5.

3. Find the displacement.
   - Record the initial position.
   - Record the final position.
4. Calculate the difference.

1. A separate numbered list after the code example follows below.

> State your sign convention before solving.

| Quantity | Symbol | Unit |
|:--|:--:|--:|
| Speed | $v=\frac{d}{t}$ | m/s |
| Acceleration | $a$ | m/s² |

```python
# Markdown and math are literal inside code
velocity = distance / time  # **not bold**, $not_math$
```

1. This list restarts at one.

\[
\begin{aligned}
v &= u + at \\
s &= ut + \frac{1}{2}at^2
\end{aligned}
\]

$$\begin{pmatrix}1 & 2 \\ 3 & 4\end{pmatrix}$$

Inline roots $\sqrt{x^2+y^2}$, indices $x_i^2$, and a vector $\vec{v}$.
'''


def sample_document(output):
    sections = [{"heading": "Formatting practice", "section_instructions": "Read the **givens** and show your work.",
        "gap": 0, "problems": [{"question": FORMATTING_SAMPLE, "answers": [5], "units": ["m/s"]}]}]
    return create_doc("Worksheet formatting review", lambda: sections, 1, output_path=output, open_document=False)


class ProgressAndFormattingTests(unittest.TestCase):
    def test_every_generator_type_has_a_unique_explicit_identity(self):
        identities = set(CUSTOM_PROGRESS)
        total = 0
        for generator in discover_generators():
            for label, metadata in generator.stored_metadata().items():
                with self.subTest(generator=type(generator).__name__, label=label):
                    identity = metadata["id"]
                    self.assertRegex(identity, r"^[a-z0-9.-]+$")
                    self.assertNotEqual(identity, label)
                    self.assertNotIn(identity, identities)
                    identities.add(identity)
                    self.assertIn(label, metadata["aliases"])
                    self.assertEqual(generator.get_problem_metadata(label)["id"], identity)
                    total += 1
        self.assertGreaterEqual(total, 70)

    def test_all_generator_histories_migrate_without_double_counting(self):
        script = '''
from tools.validate_payloads import discover_generators
from utils.ui import Interface
for index, generator in enumerate(discover_generators()):
    ui = Interface(f"history{index}", "Practice", generator, None, ["Easy", "Medium", "Hard"])
    if not ui.state.has("performance"):
        history = {}
        for label, metadata in generator.stored_metadata().items():
            history[label] = {"Medium": {"attempts": 4, "correct": 2}}
            history[metadata["id"]] = {"Medium": {"attempts": 3, "correct": 1}}
        ui.state.set("performance", history)
    ui.initialize_session_state()
    ui.performance_dropdown()
'''
        app = AppTest.from_string(script, default_timeout=25).run()
        app.run()
        self.assertFalse(app.exception)
        for index, generator in enumerate(discover_generators()):
            history = app.session_state[f"history{index}_performance"]
            self.assertEqual(set(history), {m["id"] for m in generator.stored_metadata().values()})
            self.assertTrue(all(stats["Medium"] == {"attempts": 7, "correct": 3} for stats in history.values()))

    def test_custom_activity_aliases_merge_once(self):
        for identity, metadata in CUSTOM_PROGRESS.items():
            history = {metadata["label"]: {"Hard": {"attempts": 2, "correct": 1}},
                       identity: {"Hard": {"attempts": 1, "correct": 1}}}
            migrate_custom_progress(history)
            migrate_custom_progress(history)
            self.assertEqual(history, {identity: {"Hard": {"attempts": 3, "correct": 2}}})

    def test_native_equation_structures(self):
        expressions = {r"\frac{a}{b}": "f", r"\sqrt{x}": "rad", r"x_i^2": "sSubSup",
            r"\vec{v}": "acc", r"\begin{pmatrix}1&2\\3&4\end{pmatrix}": "m",
            r"\begin{aligned}v&=u+at\\s&=ut\end{aligned}": "m"}
        for expression, tag in expressions.items():
            math = office_math(expression)
            self.assertTrue(list(math.iter(qn("m:" + tag))), expression)
            self.assertFalse(list(math.iter("group")))
            # Argument containers cannot be directly nested in another argument.
            self.assertFalse(math.xpath(".//m:e/m:e | .//m:num/m:e | .//m:den/m:e",
                namespaces={"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}))
        with self.assertRaises(ValueError):
            office_math(r"\unsupportedcommand{x}")

    def test_document_structure_survives_save_and_reopen(self):
        output = io.BytesIO()
        sample_document(output)
        output.seek(0)
        document = Document(output)
        xml = document.element
        self.assertEqual(len(document.inline_shapes), 0, "Equations should be editable Word objects")
        self.assertGreaterEqual(len(xml.xpath("//m:oMath")), 7)
        self.assertGreaterEqual(len(xml.xpath("//w:tbl")), 2)
        self.assertEqual(len(xml.xpath("//w:tblHeader")), 1)
        self.assertTrue(xml.xpath("//w:hyperlink"))
        self.assertTrue(xml.xpath("//w:strike"))
        self.assertTrue(xml.xpath("//w:br"))
        self.assertTrue(xml.xpath("//w:numPr"))
        self.assertTrue(any(r.target_ref == "https://example.com/physics" for r in document.part.rels.values()))
        runs = [r for p in document.tables[0].cell(0, 0).paragraphs for r in p.runs]
        self.assertTrue(any(r.bold and r.italic and "nested" in r.text for r in runs))
        self.assertTrue(any(r.font.name == "Consolas" and "$not_math$" in r.text for r in runs))
        numbers = document.part.numbering_part.element
        self.assertTrue(numbers.xpath('./w:abstractNum/w:lvl/w:start[@w:val="3"]'))
        self.assertTrue(numbers.xpath('./w:abstractNum/w:lvl/w:start[@w:val="1"]'))

    def test_no_network_fetch_for_markdown_images(self):
        output = io.BytesIO()
        with self.assertRaisesRegex(ValueError, "not fetched"):
            create_doc("Images", lambda: [{"heading": "Images", "gap": 0, "problems": [
                {"question": "![remote](https://example.com/image.png)", "answers": [1], "units": [""]}]}],
                1, output_path=output, open_document=False)


if __name__ == "__main__":
    if "--sample" in sys.argv:
        output = Path(sys.argv[sys.argv.index("--sample") + 1])
        output.parent.mkdir(parents=True, exist_ok=True)
        print(sample_document(output))
    else:
        unittest.main()

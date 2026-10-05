"""Worksheet content, UI lifecycle, and optional real LibreOffice integration."""
import io
from pathlib import Path
import subprocess
import sys
import unittest
import threading
import time
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from PIL import Image
from streamlit.testing.v1 import AppTest
from utils.worksheet_ui import build_worksheet, generate_sections, render_document, PREFIX
from xtrct_docs.pdf_converter import _run_converter, _conversion_slot
from xtrct_docs.pdf_converter import docx_to_pdf, libreoffice_path, PdfConversionError
from xtrct_docs.pdf_preview import page_summary, render_page


class Worksheets(unittest.TestCase):
    def test_tension_worksheet_prompt_reuses_givens(self):
        from utils.generators.forces.tension_generator import TensionGenerator
        generator = TensionGenerator()
        scenario = generator.tension_scenario("Medium")
        with patch.object(TensionGenerator, "tension_scenario", return_value=scenario):
            practice = generator.suspension_question("Medium")
            problem = generate_sections([("Tension", "Suspension", "Medium", "mixed", 1)])[0]["problems"][0]
        self.assertEqual(problem["answers"], practice["answers"])
        self.assertEqual(problem["optional_tension_diagram"], practice["diagram_data"])
        self.assertIn(f"{scenario['mass']} kg", problem["question"])
        self.assertIn(f"{scenario['theta 1']}°", problem["question"])
        self.assertIn(f"{scenario['theta 2']}°", problem["question"])
        self.assertNotIn("vault", problem["question"])
        self.assertIn("vault", practice["question"])

    def test_tension_labels_fit_print_canvas(self):
        from xtrct_docs.tension_diagram import tension_inset
        from matplotlib.backends.backend_agg import FigureCanvasAgg
        for scenario in ((1, 2), (2, 1), (1, 3), (3, 1), (2, 4), (4, 2)):
            for angles in ((1, 1), (1, 6), (6, 1), (5, 10), (45, 45), (89, 84), (84, 89), (89, 89)):
                fig = tension_inset(dict(scenario=scenario, theta1=angles[0], theta2=angles[1]))
                canvas = FigureCanvasAgg(fig)
                canvas.draw()
                for label in fig.axes[0].texts:
                    box = label.get_window_extent(canvas.get_renderer())
                    margin = .05 * fig.dpi
                    self.assertGreaterEqual(box.x0, margin)
                    self.assertGreaterEqual(box.y0, margin)
                    self.assertLessEqual(box.x1, fig.bbox.width - margin)
                    self.assertLessEqual(box.y1, fig.bbox.height - margin)
                fig.clear()

    def test_answer_boxes_are_optional(self):
        sections = generate_sections([("Tension", "Suspension", "Medium", "mixed", 3)])
        original = repr(sections)
        keys = []
        for enabled in (True, False):
            doc = Document(io.BytesIO(render_document(sections, "Compact", True,
                "Forces Practice", diagrams=True, answer_boxes=enabled)))
            xml = doc.element.xml
            self.assertEqual(xml.count("Final Answers:"), 3 if enabled else 0)
            self.assertEqual(len(doc.inline_shapes), 3)
            keys.append([p.text for p in doc.paragraphs])
            self.assertEqual(repr(sections), original)
        self.assertEqual(keys[0], keys[1])

    def test_optional_tension_inset_preserves_questions(self):
        sections = generate_sections([("Tension", "Suspension", "Medium", "mixed", 3)])
        original = repr(sections)
        for enabled, count in [(False, 0), (True, 3), (False, 0)]:
            document = Document(io.BytesIO(render_document(sections, "Compact", True, "Forces Practice", diagrams=enabled)))
            self.assertEqual(len(document.inline_shapes), count)
            self.assertEqual(repr(sections), original)
            if enabled:
                self.assertTrue(all(abs(shape.width.inches - 2) < .01 for shape in document.inline_shapes))
                self.assertTrue(all(abs(shape.height.inches - 1.05) < .01 for shape in document.inline_shapes))

    def test_answer_key_numeric_formatting(self):
        from xtrct_docs.document_creator import format_answer_key_value
        for value, expected in [(368.22809458253962, "368.228"), (2.0, "2"),
                                (-12.34567, "-12.346"), (0.125, "0.125"),
                                (0, "0"), (-0.0, "0"), (0.0014, "1.400e-03"), (0.0000123456, "1.235e-05"),
                                ("Slowing Down", "Slowing Down"), ("B", "B")]:
            self.assertEqual(format_answer_key_value(value), expected)
        original = [368.22809458253962, -12.34567]
        sections = [{"heading": "Rounding", "gap": 0, "problems": [
            {"question": "Find both tensions.", "answers": original[:], "units": ["Tension 1 (N)", "Tension 2 (N)"]},
            {"question": "Find the force.", "answers": [original[0]], "units": ["Force (N)"]}]}]
        document = Document(io.BytesIO(render_document(sections, "Compact", True)))
        key = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Tension 1 (N): 368.228", key)
        self.assertIn("Tension 2 (N): -12.346", key)
        self.assertIn("Force (N): 368.228", key)
        self.assertNotIn(str(original[0]), key)
        self.assertEqual(sections[0]["problems"][0]["answers"], original)

    def test_forces_topics_and_targets(self):
        from utils.worksheet_ui import FORCES_ACTIVITIES, activity_difficulties
        from utils.solve_for import target_options
        for activity, cls in FORCES_ACTIVITIES.items():
            for kind in cls().get_problem_types():
                for difficulty in activity_difficulties(activity):
                    for target in ["mixed"] + [o["id"] for o in target_options(cls(), kind, difficulty)]:
                        with self.subTest(activity=activity, kind=kind, difficulty=difficulty, target=target):
                            sections = generate_sections([(activity, kind, difficulty, target, 1)])
                            problem = sections[0]["problems"][0]
                            self.assertTrue(problem["answers"])
                            self.assertFalse(problem.get("diagram_data"))
                            document = Document(io.BytesIO(render_document(sections, "Compact", True, "Forces Practice")))
                            self.assertIn("Forces Practice", "\n".join(p.text for p in document.paragraphs))

    def test_forces_navigation_and_isolation(self):
        app = AppTest.from_file("Home.py", default_timeout=30).run()
        app.session_state["router_section"] = "Dynamics"
        app.session_state["router_activity"] = "Newton's Second Law"
        app.session_state[PREFIX + "count"] = 9
        app.run()
        app.button(key="worksheet_shortcut_Newton's Second Law").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(app.selectbox(key="forces_worksheet_activity").options, ["Newton's Second Law"])
        self.assertEqual(app.session_state[PREFIX + "count"], 9)
        app.session_state["nav_level"] = "advanced"
        app.run()
        self.assertIn("Atwood Machines", app.selectbox(key="forces_worksheet_activity").options)
        app.selectbox(key="forces_worksheet_activity").select("Atwood Machines").run()
        self.assertEqual(app.selectbox(key="forces_worksheet_difficulty").options, ["Medium"])
        app.button(key="forces_worksheet_add").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.session_state["forces_worksheet_section_ids"]), 2)

    def test_graph_practice_fixed_full_set(self):
        app = AppTest.from_file("Home.py", default_timeout=30).run()
        app.session_state["router_section"] = "Kinematics"
        app.session_state["router_activity"] = "Types of Motion Graphs"
        app.session_state["motion_graph_difficulty_select_unified"] = "Easy"
        app.run()
        self.assertFalse(app.exception)
        self.assertFalse(any(s.label == "Difficulty" for s in app.selectbox))
        self.assertEqual(app.session_state["motion_graph_difficulty"], "Hard")
        app.button(key="worksheet_shortcut_Types of Motion Graphs").click().run()
        self.assertFalse(any(s.label == "Worksheet difficulty" for s in app.selectbox))
        self.assertEqual(app.session_state[PREFIX + "difficulty"], "Hard")
        self.assertFalse(app.exception)

    def test_matching_choices_variety_and_persistence(self):
        from xtrct_docs.matching_graph_adapter import WorksheetMatchingGraphs
        generator = WorksheetMatchingGraphs()
        for order in generator.get_problem_types():
            payloads = [generator.choose_problem_dict(order, "Standard") for _ in range(12)]
            shapes = [p["diagram_data"]["primary"] for p in payloads]
            self.assertEqual(len(set(shapes[:6])), 6)
            self.assertEqual(len(set(shapes[6:])), 6)
            for p in payloads:
                spec = p["diagram_data"]
                self.assertEqual(len(set(spec["choices"])), 3)
                self.assertEqual(spec["choices"]["ABC".index(p["answers"][0])], spec["primary"])
            sections = generate_sections([("Matching Motion Graphs", order, "Standard", "mixed", 2)])
            before = repr([p["diagram_data"] for p in sections[0]["problems"]])
            for spacing in ("Compact", "Standard"):
                doc = Document(io.BytesIO(render_document(sections, spacing, True)))
                self.assertEqual(len(doc.inline_shapes), 2)
                self.assertEqual(repr([p["diagram_data"] for p in sections[0]["problems"]]), before)

    def test_matching_selector_and_difficulty_reset(self):
        app = AppTest.from_string("from utils.worksheet_ui import render_kinematics_worksheet\nrender_kinematics_worksheet()").run()
        app.selectbox(key=PREFIX + "activity").select("Matching Motion Graphs").run()
        self.assertEqual(app.selectbox(key=PREFIX + "difficulty").value, "Standard")
        self.assertTrue(app.selectbox(key=PREFIX + "difficulty").disabled)
        app.selectbox(key=PREFIX + "kind").select("Velocity-Time First").run()
        self.assertFalse(app.exception)
        app.selectbox(key=PREFIX + "activity").select("Constant Motion").run()
        self.assertEqual(app.selectbox(key=PREFIX + "difficulty").value, "Medium")
        self.assertFalse(app.exception)

    def test_graph_variety_and_answer_breaks(self):
        from collections import Counter
        for difficulty, pool_size in (("Easy", 6), ("Medium", 6), ("Hard", 6)):
            sections = generate_sections([("Types of Motion Graphs", "Position-Time Graph", difficulty, "mixed", 5)] * 2)
            shapes = [p["diagram_data"]["graph_type"] for s in sections for p in s["problems"]]
            for start in range(0, len(shapes), pool_size):
                chunk = shapes[start:start + pool_size]
                self.assertEqual(len(set(chunk)), len(chunk))
            counts = Counter(shapes)
            self.assertLessEqual(max(counts.values()) - min(counts.values()), 1)
            self.assertTrue(all(a != b for a, b in zip(shapes, shapes[1:])))
        doc = Document(io.BytesIO(render_document(sections, "Compact", True)))
        entries = [p for p in doc.paragraphs if "Multiple Answers:" in p.text]
        self.assertEqual(len(entries), 10)
        self.assertTrue(all("\nDirection:" in p.text and "\nMotion State:" in p.text for p in entries))

    def test_graph_shapes_answers_and_export_stability(self):
        import matplotlib.pyplot as plt
        import numpy as np
        from xtrct_docs.motion_graph_adapter import WorksheetMotionGraphs
        generator = WorksheetMotionGraphs()
        original_figures = plt.get_fignums()
        for family in generator.get_problem_types():
            for shape in generator.graph_types:
                with patch("xtrct_docs.motion_graph_adapter.random.choice", return_value=shape):
                    problem = generator.choose_problem_dict(family, "Hard")
                figure = generator.generate_diagram(problem["diagram_data"], family, "Hard")
                x, y = figure.axes[0].lines[0].get_data()
                velocity = np.diff(y) / np.diff(x) if family == "Position-Time Graph" else y[1:-1]
                self.assertEqual(problem["answers"][0], "Positive" if np.mean(velocity) > 0 else "Negative")
                speed_change = abs(velocity[-1]) - abs(velocity[0])
                expected = "Constant Velocity" if abs(speed_change) < 1e-8 else "Speeding Up" if speed_change > 0 else "Slowing Down"
                self.assertEqual(problem["answers"][1], expected)
                figure.clear()
            for difficulty, expected_count in (("Easy", 6), ("Medium", 6), ("Hard", 6)):
                with patch("xtrct_docs.motion_graph_adapter.random.choice", side_effect=lambda choices: choices[-1]) as choice:
                    WorksheetMotionGraphs().choose_problem_dict(family, difficulty)
                    self.assertEqual(len(choice.call_args.args[0]), expected_count)
        sections = generate_sections([("Types of Motion Graphs", "Position-Time Graph", "Hard", "mixed", 3)])
        specs = [dict(p["diagram_data"]) for p in sections[0]["problems"]]
        for spacing in ("Compact", "Standard", "Extra writing space"):
            doc = Document(io.BytesIO(render_document(sections, spacing, True)))
            self.assertEqual(len(doc.inline_shapes), 3)
            self.assertEqual([p["diagram_data"] for p in sections[0]["problems"]], specs)
        self.assertEqual(plt.get_fignums(), original_figures)

    def test_new_text_activities_all_targets(self):
        from utils.worksheet_ui import ACTIVITIES
        from utils.solve_for import target_options
        for activity in ("Distance & Displacement", "Projectiles", "Relative Motion"):
            generator = ACTIVITIES[activity]()
            for kind in generator.get_problem_types():
                for level in ("Easy", "Medium", "Hard"):
                    for target in ["mixed"] + [o["id"] for o in target_options(generator, kind, level)]:
                        with self.subTest(activity=activity, kind=kind, level=level, target=target):
                            sections = generate_sections([(activity, kind, level, target, 1)])
                            self.assertIsNone(sections[0]["problems"][0].get("diagram_data"))
                            document = Document(io.BytesIO(render_document(sections, "Standard", True)))
                            self.assertEqual(len(document.inline_shapes), 0)
                            self.assertNotIn("Diagram omitted", document.tables[0].cell(0, 0).text)

    def test_advanced_activity_filter(self):
        app = AppTest.from_string("from utils.worksheet_ui import render_kinematics_worksheet\nrender_kinematics_worksheet()").run()
        self.assertNotIn("Relative Motion", app.selectbox(key=PREFIX + "activity").options)
        app.session_state["nav_level"] = "advanced"
        app.run()
        app.selectbox(key=PREFIX + "activity").select("Relative Motion").run()
        self.assertFalse(app.exception)
        app.session_state["nav_level"] = "high"
        app.run()
        self.assertNotEqual(app.selectbox(key=PREFIX + "activity").value, "Relative Motion")
        self.assertFalse(app.exception)

    def test_mixed_activities_and_constant_variants(self):
        from utils.generators.kinematics.const_motion_generator import ConstantMotionGenerator
        for kind in ConstantMotionGenerator().get_problem_types():
            for difficulty in ("Easy", "Medium", "Hard"):
                sections = generate_sections([("Constant Motion", kind, difficulty, "mixed", 2),
                                              ("Accelerated Motion", "No Time", "Medium", "mixed", 2)])
                document = Document(io.BytesIO(render_document(sections, "Standard", True)))
                self.assertEqual(len(document.tables), 4)
                self.assertIn("Constant Motion", sections[0]["heading"])

    def test_activity_shortcut_routes_to_builder(self):
        app = AppTest.from_file("Home.py", default_timeout=30).run()
        app.session_state["router_section"] = "Kinematics"
        app.session_state["router_activity"] = "Accelerated Motion"
        app.run()
        self.assertFalse(any(b.key == PREFIX + "generate" for b in app.button))
        app.button(key="worksheet_shortcut_Accelerated Motion").click().run()
        self.assertEqual(app.session_state["router_activity"], "Create a worksheet")
        self.assertEqual(app.selectbox(key=PREFIX + "activity").value, "Accelerated Motion")
        app.selectbox(key=PREFIX + "kind").select("No Time").run()
        app.selectbox(key=PREFIX + "activity").select("Constant Motion").run()
        self.assertEqual(app.selectbox(key=PREFIX + "kind").value, "Constant Speed")
        self.assertFalse(app.exception)

    def test_multiple_sections_and_total_cap(self):
        sections = generate_sections([("No Time", "Easy", "mixed", 6), ("No Distance", "Hard", "mixed", 6)])
        doc = Document(io.BytesIO(render_document(sections, "Standard", True)))
        self.assertEqual(len(doc.tables), 12)
        self.assertTrue(any("Section 2: Accelerated Motion: No Distance" in p.text for p in doc.paragraphs))
        with patch("utils.worksheet_ui.DocumentGenerator.question") as question:
            with self.assertRaises(ValueError):
                generate_sections([("Mixed", "Easy", "mixed", 11)] * 2)
            question.assert_not_called()

    def test_waiting_and_busy_recovery(self):
        stages = []
        def convert(command, **kwargs):
            Path(command[-1]).with_suffix(".pdf").write_bytes(b"%PDF-fixture")
            return subprocess.CompletedProcess(command, 0)
        with patch("xtrct_docs.pdf_converter.libreoffice_path", return_value="soffice"), patch(
            "xtrct_docs.pdf_converter._run_converter", side_effect=convert
        ):
            _conversion_slot.acquire()
            try:
                with self.assertRaisesRegex(PdfConversionError, "busy"):
                    docx_to_pdf(b"docx", wait_timeout=0.01, on_stage=stages.append)
            finally:
                _conversion_slot.release()
            self.assertEqual(stages, ["Waiting for PDF conversionâ€¦"])
            _conversion_slot.acquire()
            timer = threading.Timer(0.05, _conversion_slot.release)
            timer.start()
            try:
                timings = {}
                docx_to_pdf(b"docx", on_stage=stages.append, timings=timings)
                self.assertGreater(timings["wait"], 0)
                self.assertEqual(stages[-1], "Creating your PDFâ€¦")
            finally:
                timer.join()

    def test_actual_hanging_process_is_stopped(self):
        real_popen = subprocess.Popen
        processes = []
        def record(*args, **kwargs):
            process = real_popen(*args, **kwargs)
            processes.append(process)
            return process
        with patch("xtrct_docs.pdf_converter.subprocess.Popen", side_effect=record):
            with self.assertRaises(subprocess.TimeoutExpired):
                _run_converter([sys.executable, "-c", "import time; time.sleep(60)"], timeout=0.2)
        self.assertIsNotNone(processes[0].poll())
        self.assertEqual(_run_converter([sys.executable, "-c", "print('recovered')"], timeout=5).returncode, 0)

    @patch("utils.worksheet_ui.render_page", return_value=Image.new("RGB", (10, 10), "white"))
    def test_add_remove_sections_and_limit_ui(self, renderer):
        app = AppTest.from_string("from utils.worksheet_ui import render_accelerated_worksheet\nrender_accelerated_worksheet()")
        with patch("utils.worksheet_ui.libreoffice_path", return_value="soffice"), patch(
            "utils.worksheet_ui.docx_to_pdf", return_value=b"fixture"
        ), patch("utils.worksheet_ui.page_summary", return_value={"total": 3, "questions": 2, "answers": 1}):
            app.run()
            app.button(key=PREFIX + "add").click().run()
            app.selectbox(key=PREFIX + "section_1_kind").select("No Time").run()
            app.button(key=PREFIX + "generate").click().run()
            self.assertEqual(len(app.session_state[PREFIX + "result"]["sections"]), 2)
            app.number_input(key=PREFIX + "count").set_value(20).run()
            self.assertTrue(app.button(key=PREFIX + "generate").disabled)
            self.assertFalse(app.get("download_button"))
            app.button(key=PREFIX + "remove_1").click().run()
            self.assertFalse(app.button(key=PREFIX + "generate").disabled)
            self.assertFalse(app.exception)

    def test_windows_install_discovery_without_path(self):
        with patch("xtrct_docs.pdf_converter.shutil.which", return_value=None), patch(
            "xtrct_docs.pdf_converter.sys.platform", "win32"
        ), patch("xtrct_docs.pdf_converter.Path.is_file", return_value=True):
            self.assertTrue(libreoffice_path().endswith("soffice.com"))

    def test_content_and_answer_option(self):
        for answers in (True, False):
            data = build_worksheet("No Time", "Hard", "mixed", 5, "Standard", answers)
            document = Document(io.BytesIO(data))
            self.assertEqual(len(document.tables), 5)
            self.assertEqual(any("Answer Key" in p.text for p in document.paragraphs), answers)
            if answers:
                self.assertEqual(document.sections[-1].start_type, 2)  # NEW_PAGE, not ODD_PAGE
        with self.assertRaises(ValueError):
            build_worksheet("Mixed", "Medium", "mixed", 21, "Standard", True)

    def test_conversion_failure_and_cleanup(self):
        with patch("xtrct_docs.pdf_converter.libreoffice_path", return_value=None):
            with self.assertRaises(PdfConversionError):
                docx_to_pdf(b"docx")
        directories = []
        def convert(command, **kwargs):
            source = Path(command[-1])
            directories.append(source.parent)
            self.assertEqual(source.read_bytes(), b"docx")
            source.with_suffix(".pdf").write_bytes(b"%PDF-1.7\nfixture")
            return subprocess.CompletedProcess(command, 0)
        with patch("xtrct_docs.pdf_converter.libreoffice_path", return_value="soffice"):
            with patch("xtrct_docs.pdf_converter._run_converter", return_value=subprocess.CompletedProcess([], 1)):
                with self.assertRaisesRegex(PdfConversionError, "failed"):
                    docx_to_pdf(b"docx")
            with patch("xtrct_docs.pdf_converter._run_converter", side_effect=subprocess.TimeoutExpired("soffice", 1)):
                with self.assertRaises(PdfConversionError):
                    docx_to_pdf(b"docx")
            with patch("xtrct_docs.pdf_converter._run_converter", side_effect=convert):
                self.assertTrue(docx_to_pdf(b"docx").startswith(b"%PDF-"))
                docx_to_pdf(b"docx")
        self.assertNotEqual(*directories)
        self.assertTrue(all(not path.exists() for path in directories))

    @patch("utils.worksheet_ui.render_page", return_value=Image.new("RGB", (10, 10), "white"))
    def test_page_state_and_retry(self, preview_renderer):
        app = AppTest.from_file("Home.py", default_timeout=30)
        with patch("utils.worksheet_ui.libreoffice_path", return_value="soffice"), patch(
            "utils.worksheet_ui.docx_to_pdf", side_effect=PdfConversionError("Please retry")
        ):
            app.run()
            app.session_state["router_section"] = "Kinematics"
            app.session_state["router_activity"] = "Create a worksheet"
            app.run()
            self.assertFalse(app.exception)
            app.button(key=PREFIX + "generate").click().run()
            result = app.session_state[PREFIX + "result"]
            original = result["docx"]
            self.assertEqual(result["error"], "Please retry")
            self.assertFalse(app.get("download_button"))
        with patch("utils.worksheet_ui.libreoffice_path", return_value="soffice"), patch(
            "utils.worksheet_ui.docx_to_pdf", return_value=b"%PDF-fixture"
        ) as converter, patch("utils.worksheet_ui.page_summary", return_value={"total": 2, "questions": 1, "answers": 1}):
            app.button(key=PREFIX + "retry").click().run()
            converter.assert_called_once()
            self.assertEqual(converter.call_args.args, (original,))
            app.run()
            self.assertEqual(len(app.get("download_button")), 1)
            self.assertEqual(app.get("download_button")[0].label, "Download worksheet")
            self.assertTrue(app.checkbox(key=PREFIX + "preview").value)
            preview_renderer.assert_called_once()
            self.assertEqual(app.session_state[PREFIX + "result"]["docx"], original)
            app.selectbox(key=PREFIX + "kind").select("No Time").run()
            options = app.selectbox(key=PREFIX + "target").options
            self.assertGreater(len(options), 1)
            self.assertTrue(any("Settings changed" in item.value for item in app.info))
            self.assertFalse(app.get("download_button"))
            self.assertFalse(app.exception)

    def test_no_renderer_disables_generation(self):
        app = AppTest.from_string("from utils.worksheet_ui import render_accelerated_worksheet\nrender_accelerated_worksheet()")
        with patch("utils.worksheet_ui.libreoffice_path", return_value=None):
            app.run()
            self.assertTrue(app.button(key=PREFIX + "generate").disabled)
            self.assertFalse(app.exception)
            self.assertFalse(app.get("download_button"))

    @patch("utils.worksheet_ui.render_page", return_value=Image.new("RGB", (10, 10), "white"))
    def test_layout_keeps_questions(self, preview_renderer):
        app = AppTest.from_string("from utils.worksheet_ui import render_accelerated_worksheet\nrender_accelerated_worksheet()")
        with patch("utils.worksheet_ui.libreoffice_path", return_value="soffice"), patch(
            "utils.worksheet_ui.docx_to_pdf", return_value=b"fixture"
        ), patch("utils.worksheet_ui.page_summary", return_value={"total": 2, "questions": 1, "answers": 1}):
            app.run()
            app.button(key=PREFIX + "generate").click().run()
            original = app.session_state[PREFIX + "result"]["sections"]
            app.selectbox(key=PREFIX + "spacing").select("Compact").run()
            self.assertEqual(app.button(key=PREFIX + "generate").label, "Update layout")
            app.button(key=PREFIX + "generate").click().run()
            self.assertEqual(app.session_state[PREFIX + "result"]["sections"], original)
            self.assertTrue(app.checkbox(key=PREFIX + "answer_boxes").value)
            app.checkbox(key=PREFIX + "answer_boxes").uncheck().run()
            self.assertEqual(app.button(key=PREFIX + "generate").label, "Update layout")
            app.button(key=PREFIX + "generate").click().run()
            self.assertEqual(app.session_state[PREFIX + "result"]["sections"], original)
            self.assertFalse(app.exception)

    @unittest.skipUnless(libreoffice_path(), "LibreOffice not installed")
    def test_real_pdf(self):
        data = build_worksheet("No Acceleration", "Hard", "mixed", 5, "Standard", True)
        pdf = docx_to_pdf(data)
        self.assertTrue(pdf.startswith(b"%PDF-"))
        self.assertGreater(len(pdf), 1000)
        summary = page_summary(pdf)
        self.assertEqual(summary["answers"], 1)
        self.assertGreater(summary["questions"], 0)
        self.assertTrue(render_page(pdf, 0).startswith(b"\x89PNG"))


if __name__ == "__main__":
    unittest.main()
